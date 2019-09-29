from docx import Document
from copy import copy
import xlsxwriter
import json
import os
import datetime
import re

def string_to_date(s):
    try:
        d = [int(n) for n in s.split(sep='.')[::-1]]
        return datetime.datetime(d[0], d[1], d[2])
    except:
        print('Error date: ' + s)
        return datetime.datetime(2000 ,1, 1)
    

def parse_docs_to_json():

    path = '/Users/imac/Русский язык/Отчёты о количестве учащихся/'

    files = os.listdir(path)

    students = {}
    # { name_1 : [ { ..., arrival_1 }, { ..., arrival_2 } ] ,
    #   name_2 : [ { ..., arrival_1 }, { ..., arrival_2 } ] ... }

    for f in files:
        if f.endswith('docx') & (len(f) == 15):
            print(f)
            doc = Document(path + f)
            for row in doc.tables[0].rows[1:]:
                row_cells = row.cells

                student_name = row_cells[0].text.strip().replace('ё', 'е').replace('  ', ' ')
                entry = {}
                entry['birth'] = row_cells[1].text.strip()
                entry['class'] = row_cells[2].text.strip()
                entry['city'] = row_cells[3].text.strip()
                entry['school'] = row_cells[4].text.strip()
                entry['arrival'] = row_cells[5].text.strip()

                if len(student_name) > 1:
                    if not student_name in students.keys():
                        students[student_name] = []

                    if not entry in students[student_name]:
                        students[student_name].append(entry)

    json_struct = json.dumps(students, sort_keys=True, indent=2, ensure_ascii=False)
    with open('dump_students.json','w') as json_file:
        json_file.write(json_struct)


def compile_json():
    students = {}

    with open('dump_students.json') as f:
        students = json.load(f)

    for student_name in students:
        student_entries = students[student_name]

        arrival_dates = [e['arrival'] for e in student_entries]
        latest_arrival_date = max(arrival_dates, key=string_to_date)

        latest_entry_index = arrival_dates.index(latest_arrival_date)
    
        cities = [e['city'] for e in student_entries]
        city = max(cities, key=len)

        schools = [e['school'] for e in student_entries]
        school = max(schools, key=len)
 
        student_entry_compiled = {
            "arrival": latest_arrival_date,
            "birth": student_entries[latest_entry_index]['birth'],  # take latest
            "city": city,
            "class": student_entries[latest_entry_index]['class'],
            "school": school
            }

        students[student_name] = student_entry_compiled

    json_struct = json.dumps(students, sort_keys=True, indent=2, ensure_ascii=False)
    with open('compiled_students.json','w') as json_file:
        json_file.write(json_struct)    
        
    with open('students.txt','w') as f:
        for student_name in students:
            f.write(student_name + '\n')


def convert_json_to_xls():

    with open('compiled_students.json') as f:
        students = json.load(f)


    workbook = xlsxwriter.Workbook('students.xlsx')
    
    worksheet = workbook.add_worksheet()



    for i, student in enumerate(students):
        worksheet.write(i, 0, student)
        worksheet.write(i, 1, students[student]['birth'])
        worksheet.write(i, 2, students[student]['class'])
        worksheet.write(i, 3, students[student]['city'])
        worksheet.write(i, 4, students[student]['school'])
        worksheet.write(i, 5, students[student]['arrival'])

    fmt = {'border': 1, 'font_size': 12}
    cell_format = workbook.add_format(fmt)
    fmt['align'] = 'center'
    cell_format_centered = workbook.add_format(fmt)

    worksheet.set_column(0, 0, 32, cell_format)
    worksheet.set_column(1, 1, 10, cell_format)
    worksheet.set_column(2, 2, 6,  cell_format_centered)
    worksheet.set_column(3, 3, 22, cell_format)
    worksheet.set_column(4, 4, 22, cell_format)
    worksheet.set_column(5, 5, 10, cell_format)

    workbook.close()

        
parse_docs_to_json()
compile_json()
convert_json_to_xls()

